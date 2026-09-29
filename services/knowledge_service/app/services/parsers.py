from dataclasses import asdict, dataclass
from io import BytesIO
from pathlib import Path

from docx import Document
from pptx import Presentation
from pypdf import PdfReader


@dataclass
class ParsedUnit:
    text: str
    page_no: int | None = None
    section: str | None = None

    def to_dict(self) -> dict:
        return asdict(self)


def parse_pdf(data: bytes) -> list[ParsedUnit]:
    reader = PdfReader(BytesIO(data))

    units = []

    for index, page in enumerate(
        reader.pages,
        start=1,
    ):
        text = page.extract_text() or ""

        text = text.strip()

        if text:
            units.append(
                ParsedUnit(
                    text=text,
                    page_no=index,
                    section=f"page-{index}",
                )
            )

    return units


def parse_pptx(data: bytes) -> list[ParsedUnit]:
    presentation = Presentation(
        BytesIO(data)
    )

    units = []

    for index, slide in enumerate(
        presentation.slides,
        start=1,
    ):
        texts = []

        for shape in slide.shapes:
            if hasattr(shape, "text"):
                value = shape.text.strip()

                if value:
                    texts.append(value)

        text = "\n".join(texts).strip()

        if text:
            units.append(
                ParsedUnit(
                    text=text,
                    page_no=index,
                    section=f"slide-{index}",
                )
            )

    return units


def parse_docx(data: bytes) -> list[ParsedUnit]:
    document = Document(
        BytesIO(data)
    )

    texts = [
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]

    text = "\n".join(texts)

    if not text:
        return []

    return [
        ParsedUnit(
            text=text,
            section="document",
        )
    ]


def parse_text(data: bytes) -> list[ParsedUnit]:
    text = data.decode(
        "utf-8",
        errors="replace",
    ).strip()

    if not text:
        return []

    return [
        ParsedUnit(
            text=text,
            section="document",
        )
    ]


def parse_material(
    *,
    filename: str,
    data: bytes,
) -> tuple[str, list[ParsedUnit]]:
    suffix = Path(filename).suffix.lower()

    if suffix == ".pdf":
        return "pypdf", parse_pdf(data)

    if suffix == ".pptx":
        return "python-pptx", parse_pptx(data)

    if suffix == ".docx":
        return "python-docx", parse_docx(data)

    if suffix in {
        ".txt",
        ".md",
        ".markdown",
    }:
        return "text", parse_text(data)

    raise ValueError(
        f"Unsupported parser for file type: "
        f"{suffix or '<none>'}"
    )
